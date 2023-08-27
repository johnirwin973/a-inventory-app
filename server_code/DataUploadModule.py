import anvil.email
import anvil.media
import anvil.users
import anvil.tables as tables
import anvil.tables.query as q
from anvil.tables import app_tables
import anvil.server
import openpyxl


@anvil.server.callable
def import_csv(file):
    with anvil.media.TempFile(file) as temp_file:
        with open(temp_file, "r", encoding='latin-1') as f:
            csv_data = f.readlines()
            rows = [row.strip().split(",") for row in csv_data]
            column_headers = rows[0]
            Description = "Description"
            item_number_column = "Item_number"  
            cost_column = "Cost" 
            image_column = "Image"  
            for row in rows[1:]:
                row_data = dict(zip(column_headers, row))
                item_number = row_data[item_number_column]
                existing_row = app_tables.inventory.get(Item_number=item_number)
                if existing_row is not None:
                   
                    if existing_row[cost_column] != row_data[cost_column]:
                        existing_row[cost_column] = row_data[cost_column]
                    
                    if image_column in existing_row:
                        row_data[image_column] = existing_row[image_column]
                    existing_row.update()
                else:
                    
                    app_tables.inventory.add_row(**row_data)



@anvil.server.callable
def update_prices_quantity(file):
    with anvil.media.TempFile(file) as temp_file:
        with open(temp_file, "r", encoding='latin-1') as f:
            csv_data = f.readlines()
            rows = [row.strip().split(",") for row in csv_data]
            column_headers = rows[0]
            item_number_column = "Item_number"
            cost_column = "Cost"
            quantity_column = "Quantity"
            stock_limit_column = "stock_limit"  
            
            def preprocess_cost(cost_str):
                return cost_str.replace('$', '')  
            
            def preprocess_quantity(quantity_str):
                return int(quantity_str) if quantity_str else 0  
            
            def preprocess_stock_limit(limit_str):
                return int(limit_str) if limit_str else 0  
            
            data_dict = {row_data[column_headers.index(item_number_column)]: {
                cost_column: preprocess_cost(row_data[column_headers.index(cost_column)]),
                quantity_column: preprocess_quantity(row_data[column_headers.index(quantity_column)]),
                stock_limit_column: preprocess_stock_limit(row_data[column_headers.index(stock_limit_column)])
            } for row_data in rows[1:]}
            
            inventory_table = app_tables.inventory.search()
            for row in inventory_table:
                item_number = row[item_number_column]
                if item_number in data_dict:
                    new_cost = data_dict[item_number][cost_column]
                    new_quantity = data_dict[item_number][quantity_column]
                    new_stock_limit = data_dict[item_number][stock_limit_column]
                  
                    new_total =  "${:.2f}".format(float(new_cost) * new_quantity)
            
                    row[cost_column] = str(new_cost)
                    row[quantity_column] = new_quantity
                    row[stock_limit_column] = new_stock_limit
                    row["Total"] = new_total
                    row.update()

    return "Update completed"


@anvil.server.callable
def add_inventory_item(category, description, item_number, cost, image):
    app_tables.inventory.add_row(Category=category, Description=description, Item_number=item_number, Cost=cost, Images=image)

@anvil.server.callable
def update_inventory_image(item_number, new_image):
    item = app_tables.inventory.get(Item_number=item_number)
    if item:
        item['Images'] = new_image
        item.update()
        return True
    else:
        return False

@anvil.server.callable
def delete_item_row(item_number):
    # Search for the row with the given item_number
    rows = app_tables.inventory.search(Item_number=item_number)
    
    # Check if the row exists
    if len(rows) > 0:
        # Get the first row (assuming there's only one row with the given item_number)
        row = rows[0]
        # Delete the row from the data table
        row.delete()
    else:
        # Row with the given item_number not found
        raise Exception("Item not found in the inventory.")
  